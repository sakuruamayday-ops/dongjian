"""Real file regression for independently generated, unbranded documents."""
import hashlib
import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

import pymupdf as fitz
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from PIL import Image as Raster

SOURCE = Path(__file__).resolve().parents[1] / "skill-source"
RUNTIME = SOURCE / "_runtime/gongchuang-branding"
sys.path.insert(0, str(RUNTIME / "scripts"))
from neutral_document import neutralize_office, legacy_asset_hashes
from delivery_gate import validate_artifact, GateFailure
from html_branding import brand_html_text
from pdf_two_pass import brand_pdf_bytes


class NeutralDocuments(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix="dongjian-document-test-"))
        self.brand = RUNTIME / "assets/brand-red-07.png"
        self.user_image = self.root / "user-image.png"
        Raster.new("RGB", (12, 8), (27, 82, 149)).save(self.user_image)

    def assert_no_brand_media(self, path):
        hashes = legacy_asset_hashes()
        with ZipFile(path) as package:
            media = [package.read(name) for name in package.namelist() if "/media/" in name]
        self.assertTrue(media, "User image must remain")
        self.assertTrue(all(hashlib.sha256(value).hexdigest() not in hashes for value in media))

    def test_docx_retains_body_and_user_image(self):
        path = self.root / "report.docx"
        doc = Document()
        doc.add_paragraph("企业材料正文，金额 12345 元。")
        doc.add_picture(str(self.user_image))
        header = doc.sections[0].header.paragraphs[0]
        header.text = "共创研究院"
        header.add_run().add_picture(str(self.brand))
        doc.save(path)
        with self.assertRaises(GateFailure):
            validate_artifact(path)
        neutralize_office(path)
        self.assertEqual(validate_artifact(path)["watermarks"], 0)
        self.assertEqual(Document(path).paragraphs[0].text, "企业材料正文，金额 12345 元。")
        self.assert_no_brand_media(path)

    def test_xlsx_retains_formula_cells_and_user_image(self):
        path = self.root / "report.xlsx"
        book = Workbook()
        sheet = book.active
        sheet["A1"] = 12345
        sheet["A2"] = "=A1*2"
        sheet.add_image(Image(str(self.brand)), "B1")
        sheet.add_image(Image(str(self.user_image)), "D1")
        sheet.oddHeader.right.text = "共创研究院"
        book.save(path)
        neutralize_office(path)
        self.assertEqual(validate_artifact(path)["watermarks"], 0)
        result = load_workbook(path)
        self.assertEqual(result.active["A2"].value, "=A1*2")
        self.assertEqual(len(result.active._images), 1)
        result.close()
        self.assert_no_brand_media(path)

    def test_html_preserves_content_and_removes_only_known_chrome(self):
        source = '<html><head><style id="gongchuang-public-brand-style">body{color:red}</style></head><body><div class="gongchuang-document-header">共创研究院</div><p>材料正文</p></body></html>'
        path = self.root / "report.html"
        path.write_text(brand_html_text(source), encoding="utf-8")
        self.assertEqual(validate_artifact(path)["watermarks"], 0)
        self.assertIn("<p>材料正文</p>", path.read_text())

    def test_pdf_has_no_added_images_and_enforces_page_count(self):
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Source content 12345")
        data = doc.tobytes()
        doc.close()
        path = self.root / "report.pdf"
        audit = brand_pdf_bytes(data, path)
        self.assertEqual(audit, [{"page": 1, "watermarks": 0}])
        self.assertEqual(path.read_bytes(), data)
        self.assertEqual(validate_artifact(path, expected_pages=1)["pages"], 1)
        with self.assertRaises(GateFailure):
            validate_artifact(path, expected_pages=2)

    def test_pdf_rejects_an_existing_publisher_watermark(self):
        doc = fitz.open()
        page = doc.new_page()
        page.insert_image(fitz.Rect(100, 100, 300, 300), filename=str(self.brand))
        path = self.root / "branded.pdf"
        doc.save(path)
        doc.close()
        with self.assertRaises(GateFailure):
            validate_artifact(path)

    def test_bundled_templates_have_no_publisher_text_or_images(self):
        templates = list((SOURCE / "project-feasibility/assets/report-templates").rglob("*.docx"))
        self.assertTrue(templates)
        for path in templates:
            with self.subTest(template=str(path.relative_to(SOURCE))):
                self.assertEqual(validate_artifact(path)["watermarks"], 0)
                with ZipFile(path) as package:
                    for name in package.namelist():
                        if name.endswith(".xml"):
                            self.assertNotIn("共创", package.read(name).decode("utf-8"))

    def test_native_docx_generator_has_no_default_publisher_metadata(self):
        script = SOURCE / "evidence-ledger/scripts/create_docx_from_text.py"
        spec = importlib.util.spec_from_file_location("native_docx_generator", script)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        output = self.root / "generated.docx"
        module.build_document("# 企业资料核对\n营收 12345 元。", output)
        self.assertEqual(validate_artifact(output)["watermarks"], 0)
        document = Document(output)
        self.assertEqual(document.core_properties.author, "")
        self.assertIn("12345", "".join(p.text for p in document.paragraphs))


if __name__ == "__main__":
    unittest.main()
