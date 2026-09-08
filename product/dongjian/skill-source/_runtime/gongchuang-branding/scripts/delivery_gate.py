#!/usr/bin/env python3
"""Portable delivery gate for the shared Gongchuang branding runtime."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pymupdf as fitz

from brand_config import public_identity
from neutral_document import legacy_asset_hashes, office_parts


RUNTIME_ROOT = Path(__file__).resolve().parent.parent
ASSET_DIR = RUNTIME_ROOT / "assets"
WATERMARK_NAME = "Gongchuang Institute Centered Watermark v4"
XLSX_MARKER = "_GONGCHUANG_INSTITUTE_UNIFORM_WATERMARK_V4"


class GateFailure(RuntimeError):
    pass


def _asset_alpha_hashes() -> set[str]:
    hashes: set[str] = set()
    assets = sorted(ASSET_DIR.glob("brand-*.png"))
    if not assets:
        raise GateFailure(f"品牌资产缺失：{ASSET_DIR}")
    for path in assets:
        pixmap = fitz.Pixmap(str(path))
        if pixmap.alpha:
            alpha = bytes(pixmap.samples[pixmap.n - 1 :: pixmap.n])
            hashes.add(hashlib.sha256(alpha).hexdigest())
    if not hashes:
        raise GateFailure("品牌资产未提供可验证的透明度通道")
    return hashes


def pdf_brand_watermark_rects(document: fitz.Document, page: fitz.Page) -> list[fitz.Rect]:
    """Return actual placements without recounting shared image xrefs."""
    alpha_hashes = _asset_alpha_hashes()
    brand_xrefs: set[int] = set()
    for image in page.get_images(full=True):
        xref, soft_mask = int(image[0]), int(image[1])
        if xref in brand_xrefs or not soft_mask:
            continue
        alpha = fitz.Pixmap(document, soft_mask)
        if hashlib.sha256(alpha.samples).hexdigest() in alpha_hashes:
            brand_xrefs.add(xref)
    marks: list[fitz.Rect] = []
    for xref in sorted(brand_xrefs):
        marks.extend(fitz.Rect(rect) for rect in page.get_image_rects(xref))
    return marks


def validate_pdf(
    path: str | Path,
    *,
    expected_pages: int | None = None,
    expected_author: str | None = None,
    expected_title_contains: str | None = None,
) -> dict[str, Any]:
    pdf_path = Path(path)
    if not pdf_path.is_file():
        raise FileNotFoundError(pdf_path)
    document = fitz.open(pdf_path)
    try:
        if not document.page_count:
            raise GateFailure("PDF没有页面")
        if expected_pages is not None and document.page_count != expected_pages:
            raise GateFailure(f"PDF页数为{document.page_count}，要求为{expected_pages}")
        metadata = document.metadata or {}
        if expected_author and metadata.get("author") != expected_author:
            raise GateFailure(f"PDF作者元数据为{metadata.get('author')!r}，要求为{expected_author!r}")
        if expected_title_contains and expected_title_contains not in (metadata.get("title") or ""):
            raise GateFailure(f"PDF标题元数据未包含{expected_title_contains!r}")

        page_audit: list[dict[str, Any]] = []
        for page_number, page in enumerate(document, start=1):
            marks = pdf_brand_watermark_rects(document, page)
            if marks:
                raise GateFailure(f"PDF第{page_number}页仍有旧产品水印")
            if public_identity()["document_header"] in page.get_text():
                raise GateFailure(f"PDF第{page_number}页仍有旧产品页眉")
            page_audit.append({"page": page_number, "watermarks": 0})
        return {
            "status": "passed",
            "path": str(pdf_path),
            "format": "pdf",
            "pages": document.page_count,
            "watermarks": 0,
            "watermark_size": [],
            "metadata": {
                "title": metadata.get("title", ""),
                "author": metadata.get("author", ""),
                "producer": metadata.get("producer", ""),
            },
            "page_audit": page_audit,
        }
    finally:
        document.close()


def _neutral_office(path: Path) -> None:
    parts = office_parts(path)
    hashes = legacy_asset_hashes()
    for name, value in parts.items():
        if "/media/" in name and hashlib.sha256(value).hexdigest() in hashes:
            raise GateFailure(f"Office文件仍有旧产品水印资产：{name}")
        if name.endswith(".xml"):
            text = value.decode("utf-8")
            if WATERMARK_NAME in text or XLSX_MARKER in text:
                raise GateFailure(f"Office文件仍有旧产品水印标记：{name}")
            if ("header" in name.lower() or "footer" in name.lower()) and public_identity()["document_header"] in text:
                raise GateFailure(f"Office文件仍有旧产品页眉：{name}")


def validate_docx(path: str | Path) -> dict[str, Any]:
    from docx import Document
    artifact = Path(path)
    _neutral_office(artifact)
    document = Document(artifact)
    if not document.sections:
        raise GateFailure("Word没有有效章节")
    return {"status": "passed", "path": str(artifact), "format": "docx", "watermarks": 0}


def validate_xlsx(path: str | Path) -> dict[str, Any]:
    from openpyxl import load_workbook
    artifact = Path(path)
    _neutral_office(artifact)
    workbook = load_workbook(artifact, keep_vba=artifact.suffix.lower() == ".xlsm")
    try:
        visible = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
        if not visible:
            raise GateFailure("Excel没有可见工作表")
        for sheet in visible:
            for header in (sheet.oddHeader, sheet.evenHeader, sheet.firstHeader):
                if any(public_identity()["document_header"] in (part.text or "") for part in (header.left, header.center, header.right)):
                    raise GateFailure(f"Excel工作表{sheet.title}仍有旧产品页眉")
        return {"status": "passed", "path": str(artifact), "format": artifact.suffix.lower().lstrip("."), "worksheets": len(visible), "watermarks": 0}
    finally:
        workbook.close()


def validate_html(path: str | Path) -> dict[str, Any]:
    import base64
    artifact = Path(path)
    source = artifact.read_text(encoding="utf-8")
    if any(marker in source for marker in ("gongchuang-public-brand-style", "gongchuang-document-header", "gongchuang-cover-signature")):
        raise GateFailure("HTML仍有旧产品品牌样式或页眉")
    hashes = legacy_asset_hashes()
    for encoded in re.findall(r"data:image/[^;,]+;base64,([A-Za-z0-9+/=]+)", source):
        if hashlib.sha256(base64.b64decode(encoded, validate=True)).hexdigest() in hashes:
            raise GateFailure("HTML仍有旧产品水印资产")
    return {"status": "passed", "path": str(artifact), "format": "html", "watermarks": 0}


def validate_artifact(
    path: str | Path,
    *,
    check_stamp: bool = True,
    expected_pages: int | None = None,
    expected_author: str | None = None,
    expected_title_contains: str | None = None,
) -> dict[str, Any]:
    del check_stamp
    artifact = Path(path)
    suffix = artifact.suffix.lower()
    if suffix == ".pdf":
        return validate_pdf(
            artifact,
            expected_pages=expected_pages,
            expected_author=expected_author,
            expected_title_contains=expected_title_contains,
        )
    if suffix == ".docx":
        return validate_docx(artifact)
    if suffix in {".xlsx", ".xlsm"}:
        return validate_xlsx(artifact)
    if suffix in {".html", ".htm"}:
        return validate_html(artifact)
    raise GateFailure(f"便携品牌运行时不支持该格式：{suffix}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("artifact", type=Path)
    parser.add_argument("--expected-pages", type=int)
    parser.add_argument("--expected-author")
    parser.add_argument("--expected-title-contains")
    parser.add_argument("--audit-json", type=Path)
    args = parser.parse_args()
    try:
        result = validate_artifact(
            args.artifact,
            expected_pages=args.expected_pages,
            expected_author=args.expected_author,
            expected_title_contains=args.expected_title_contains,
        )
    except Exception as exc:
        print(json.dumps({"status": "blocked", "error": str(exc)}, ensure_ascii=False, indent=2))
        return 2
    if args.audit_json:
        args.audit_json.parent.mkdir(parents=True, exist_ok=True)
        args.audit_json.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
